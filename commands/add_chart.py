import command_system
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

def chart():
    wb = load_workbook('/home/votebot/mysite/commands/voting.xlsx')
    ws = wb.active
    #pie = PieChart()
    chart1 = BarChart()
    chart1.type = "col"
    chart1.style = 10
    chart1.title = "Результаты зрительского голосования"
    chart1.y_axis.title = ''
    chart1.x_axis.title = ''
    labels = Reference(ws, min_col=1, min_row=2, max_row=11)
    data = Reference(ws, min_col=2, min_row=2, max_row=11)
    #pie.add_data(data, titles_from_data=True)
    #pie.set_categories(labels)
    #pie.title = "Результаты"
    #ws.add_chart(pie, "D1")
    chart1.add_data(data, titles_from_data=False)
    chart1.set_categories(labels)
    chart1.shape = 4
    ws.add_chart(chart1, "D1")
    wb.save('/home/votebot/mysite/commands/voting.xlsx')
    return 'chart added', '', None

chart_command = command_system.Command()

chart_command.keys = ['chart']
chart_command.description = 'adds chart'
chart_command.process = chart